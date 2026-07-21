import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

def main():
    doc = docx.Document()
    
    # Setup styles
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)

    # 1. TITLE PAGE
    doc.add_paragraph("\n\n\n\n\n")
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run("FinOps Intelligence Platform\nA Custom AI-Based Solution for Multi-Cloud Financial Operations\n\n")
    title_run.bold = True
    title_run.font.size = Pt(24)

    doc.add_paragraph("\n\n\n")
    sub_para = doc.add_paragraph()
    sub_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_para.add_run("Subject: AI Lab Course Project (6th Sem)\nLanguage: Python\n\n\n")
    sub_run.font.size = Pt(14)
    
    doc.add_paragraph("\n\n\n")
    info_para = doc.add_paragraph()
    info_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info_run = info_para.add_run("Submitted by: [Student Name], [Roll No], [Department]\n")
    info_run.font.size = Pt(14)
    info_run2 = info_para.add_run("Submitted to: [Professor Name]\nInstitution: [College Name]\nDate: April 2026")
    info_run2.font.size = Pt(14)
    
    doc.add_page_break()

    def add_heading(text, level=1):
        h = doc.add_heading(text, level=level)
        for run in h.runs:
            run.font.name = 'Times New Roman'
            run.bold = True
            if level == 1:
                run.font.size = Pt(16)
                run.font.color.rgb = RGBColor(0, 51, 102)
            elif level == 2:
                run.font.size = Pt(14)
                run.font.color.rgb = RGBColor(0, 76, 153)
        return h

    # 2. ABSTRACT
    add_heading('2. Abstract', level=1)
    doc.add_paragraph("The FinOps Intelligence Platform is a comprehensive, multi-cloud financial operations solution designed to tackle the growing complexity of cloud cost management. Leveraging advanced artificial intelligence techniques, this platform provides automated detection of billing anomalies, probabilistic forecasting of future cloud costs, and an interactive, conversational AI-driven optimization interface. By analyzing multi-cloud billing data across AWS, Azure, and Google Cloud Platform (GCP), the system can proactively identify unexpected billing spikes and gradual cost drifts before they result in significant budget overruns. The platform integrates a hybrid anomaly detection ensemble using Z-score analysis, STL decomposition, and Isolation Forests, coupled with SHAP-based root cause attribution. Furthermore, cost forecasting is achieved through a Prophet and LightGBM ensemble, providing P10, P50, and P90 confidence bands. An integrated LLM chatbot, powered by the Llama 3.3-70B model via the Groq API, allows users to query live financial data and explore what-if scenarios in natural language. This report details the methodology, dataset, system architecture, and evaluation of the platform, demonstrating a robust AI-driven approach to solving real-world FinOps challenges.")
    
    # 3. OBJECTIVE
    add_heading('3. Objective', level=1)
    doc.add_paragraph("The primary objectives of this 6th Semester AI Lab Course Project are as follows:")
    doc.add_paragraph("• Develop an AI-driven system solving a real-world problem: The project aims to address the critical business challenge of cloud cost management (FinOps), an area where manual monitoring is increasingly unfeasible due to scale.", style='List Bullet')
    doc.add_paragraph("• Implement advanced AI techniques: The system must successfully integrate anomaly detection, time-series forecasting, and natural language processing (NLP) to create a holistic AI solution.", style='List Bullet')
    doc.add_paragraph("• Utilize Python for end-to-end implementation: Ensure that the data generation, backend API, machine learning pipelines, and model evaluation are robustly implemented using Python.", style='List Bullet')
    doc.add_paragraph("• Employ a realistic multi-cloud dataset: Train and evaluate the models using a comprehensive synthetic dataset that accurately mirrors the complex billing structures of AWS, Azure, and GCP.", style='List Bullet')

    # 4. PROBLEM STATEMENT
    add_heading('4. Problem Statement', level=1)
    doc.add_paragraph("In the modern enterprise IT landscape, organizations increasingly adopt multi-cloud architectures to avoid vendor lock-in and leverage the specific strengths of providers such as AWS, Azure, and GCP. However, this diversification introduces immense complexity in cost management. Cloud billing data is often highly granular, generated at massive volumes, and presented in varying formats across different providers. As a result, engineering and finance teams struggle to gain unified visibility into their cloud spend.\n\nTraditional rule-based threshold alerts are insufficient; they often trigger false positives or miss subtle, gradual increases in cost. Teams frequently face unexpected billing spikes—such as a runaway Kubernetes cluster on GCP or unoptimized bandwidth egress on Azure—which are only discovered at the end of the billing cycle, leading to severe budget overruns. Furthermore, existing tools lack actionable insights, leaving teams without clear explanations for why costs spiked or how to optimize them. \n\nThis project addresses these challenges by building an AI-driven FinOps Intelligence Platform that automatically ingests cross-cloud billing data, detects anomalies with high precision, forecasts future costs to warn of budget breaches, and provides a natural language interface for conversational insights.")

    # 5. DATASET
    add_heading('5. Dataset', level=1)
    doc.add_paragraph("Given the proprietary and sensitive nature of real-world corporate cloud billing data, this project utilizes a custom-generated, highly realistic synthetic multi-cloud billing dataset. The dataset was generated using a custom Python script (data_generator.py) designed to simulate the usage patterns and billing complexities of a mid-to-large enterprise operating across three major cloud providers: AWS, Azure, and GCP.")
    doc.add_paragraph("The generated dataset spans several months of daily aggregated service costs, incorporating natural seasonal variations (e.g., higher compute usage during weekdays, lower on weekends) and inherent noise. The data is structured with key features including Date, CloudProvider, Service, AccountID, Region, and Cost.")
    
    add_heading('Ground-Truth Anomaly Scenarios', level=2)
    doc.add_paragraph("To rigorously evaluate the anomaly detection model, the dataset includes six carefully engineered ground-truth anomaly scenarios, representing common FinOps incidents:")
    doc.add_paragraph("1. ANO-001: AWS EC2 Autoscaling Spike (4.8× increase) - Simulates a misconfigured auto-scaling group responding to a brief traffic surge but failing to scale down.", style='List Bullet')
    doc.add_paragraph("2. ANO-002: AWS S3 Gradual Drift (3.2× increase) - Simulates unchecked log accumulation and lack of lifecycle policies leading to steadily rising storage costs.", style='List Bullet')
    doc.add_paragraph("3. ANO-003: Azure Virtual Machines Weekend Spike (3.1× increase) - Represents development environments accidentally left running over the weekend.", style='List Bullet')
    doc.add_paragraph("4. ANO-004: Azure Bandwidth Egress Spike (5.5× increase) - Simulates a sudden data exfiltration event or a misconfigured CDN pulling raw origin data excessively.", style='List Bullet')
    doc.add_paragraph("5. ANO-005: GCP BigQuery Runaway Scan (6.2× increase) - Represents an unoptimized, unbounded cross-join SQL query executed by a data analyst.", style='List Bullet')
    doc.add_paragraph("6. ANO-006: GCP Kubernetes Gradual Drift (2.4× increase) - Simulates the gradual addition of microservices and inefficient bin-packing over time.", style='List Bullet')
    doc.add_paragraph("The dataset is persistently stored using an SQLite database configured in WAL (Write-Ahead Logging) mode to support concurrent reads during real-time data ingestion.")

    # 6. AI MODEL DEVELOPMENT
    add_heading('6. AI Model Development', level=1)
    
    add_heading('a) Anomaly Detection', level=2)
    doc.add_paragraph("The anomaly detection module (anomaly_detector.py) employs a hybrid ensemble approach to maximize recall while maintaining precision. Because cloud costs exhibit strong weekly seasonality, a simple static threshold is inadequate. The ensemble consists of:")
    doc.add_paragraph("• Z-Score Analysis: A statistical baseline that identifies data points lying several standard deviations away from the rolling mean.", style='List Bullet')
    doc.add_paragraph("• STL Decomposition: Seasonal and Trend decomposition using Loess (STL) is applied to extract the underlying trend and seasonality. Anomalies are flagged based on the magnitude of the residual component.", style='List Bullet')
    doc.add_paragraph("• Isolation Forest: A machine learning approach utilizing scikit-learn's IsolationForest algorithm. This tree-based model isolates anomalies by recursively partitioning the dataset; anomalies, being few and distinct, require fewer partitions to isolate.", style='List Bullet')
    doc.add_paragraph("The system uses a 'union voting' ensemble mechanism, meaning if any of the underlying models flag a point with high confidence, it is marked as an anomaly. This prioritizes maximum recall, which is critical in FinOps to ensure no cost spikes go unnoticed. Furthermore, SHAP (SHapley Additive exPlanations) is integrated to provide feature-level root cause attribution, explaining exactly which dimensions (e.g., Region, Service) contributed to the anomaly.")

    add_heading('b) Cost Forecasting', level=2)
    doc.add_paragraph("The forecasting engine (forecasting_engine.py) predicts future cloud spend to provide early warnings of budget breaches. It utilizes a robust ensemble of:")
    doc.add_paragraph("• Facebook Prophet: Excel at handling time-series data with strong seasonal effects and missing data.", style='List Bullet')
    doc.add_paragraph("• LightGBM: A gradient boosting framework that captures complex non-linear relationships and interactions between features like Service and Region.", style='List Bullet')
    doc.add_paragraph("Instead of recursive forecasting, the model uses a direct multi-step forecasting strategy to prevent the accumulation of prediction errors over long horizons. Crucially, the engine provides probabilistic predictions, outputting P10, P50, and P90 confidence bands. This allows financial stakeholders to understand the best-case, expected, and worst-case spend scenarios.")

    add_heading('c) AI Chatbot', level=2)
    doc.add_paragraph("To democratize access to financial data, an AI Chatbot (chatbot.py) is integrated. The chatbot is powered by the Groq API, utilizing the Llama 3.3-70B-Versatile model. The chatbot is context-aware: it intercepts user queries, dynamically translates them into SQL to query the live FinOps SQLite database, and synthesizes the retrieved data into natural language responses. This enables users to perform complex what-if scenarios (e.g., 'What if we reduce Azure VM usage by 20%?') and query root causes ('Why did AWS EC2 costs spike last week?') without needing to write database queries.")

    # 7. TECH STACK
    add_heading('7. Tech Stack', level=1)
    doc.add_paragraph("The project is built on a modern, robust technology stack, detailed in the table below:")
    
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Component'
    hdr_cells[1].text = 'Technology / Framework'
    
    tech_stack = [
        ("Frontend", "React 18, TypeScript, Vite, Recharts, Lucide Icons"),
        ("Backend API", "Python 3, FastAPI, Uvicorn, Pydantic"),
        ("ML / AI Pipeline", "scikit-learn (Isolation Forest), statsmodels (STL), Facebook Prophet, LightGBM, SHAP"),
        ("Large Language Model (LLM)", "Groq API → Llama 3.3-70B-Versatile"),
        ("Database", "SQLite with WAL (Write-Ahead Logging) mode"),
        ("Data Processing", "pandas, NumPy")
    ]
    
    for comp, tech in tech_stack:
        row_cells = table.add_row().cells
        row_cells[0].text = comp
        row_cells[1].text = tech

    # 8. DATA VISUALIZATION & EXPLORATION
    add_heading('8. Data Visualization & Exploration', level=1)
    doc.add_paragraph("Effective FinOps relies heavily on data visibility. The platform includes a comprehensive React-based frontend that serves as a Command Center. It utilizes Recharts for rendering interactive, responsive dashboards.")
    doc.add_paragraph("The visualization suite includes 8 distinct categories:")
    doc.add_paragraph("1. Time-Series Graphs: Tracking daily spend over time with clear visual highlights for detected anomalies.", style='List Bullet')
    doc.add_paragraph("2. Spend Distribution: Pie charts showing the proportion of spend across AWS, Azure, and GCP.", style='List Bullet')
    doc.add_paragraph("3. Service Spend: Bar charts breaking down costs by individual cloud services (e.g., EC2, S3, BigQuery).", style='List Bullet')
    doc.add_paragraph("4. Anomaly Counts: Visual tracking of anomaly frequency over time.", style='List Bullet')
    doc.add_paragraph("5. Box Plots: Illustrating the spread and outliers in spend for different services.", style='List Bullet')
    doc.add_paragraph("6. Stacked Bar Charts: Providing cross-sectional views of spend by region and provider.", style='List Bullet')
    doc.add_paragraph("7. Budget Forecasts: Line charts displaying the P10, P50, and P90 forecast bands against the allocated budget threshold.", style='List Bullet')
    doc.add_paragraph("8. Detailed Anomaly Tables: Providing tabular data with SHAP-based root cause explanations and cost impact calculations.", style='List Bullet')
    doc.add_paragraph("Additionally, a standalone Plotly HTML dashboard generator (dashboard_generator.py) is provided for offline analysis.")

    # 9. MODEL TRAINING & EVALUATION
    add_heading('9. Model Training & Evaluation', level=1)
    doc.add_paragraph("The performance of the AI models was rigorously evaluated using standard machine learning metrics.")
    doc.add_paragraph("For Anomaly Detection, the primary metrics were Precision, Recall, and F1-Score. The system was validated against the 6 engineered ground-truth anomaly labels (ANO-001 to ANO-006). Given the FinOps context, the model was tuned to prioritize Recall, ensuring that no significant cost spikes are missed, even if it results in a slightly higher false-positive rate. The real-time detection rate is dynamically calculated and displayed on the live data upload dashboard.")
    doc.add_paragraph("For Cost Forecasting, performance was measured using Mean Absolute Error (MAE) and Root Mean Squared Error (RMSE). The use of probabilistic bands (P10/P90) ensures that the model's uncertainty is accurately quantified, which is more useful for budget planning than point predictions alone.")

    # 10. HANDLING OVERFITTING & UNDERFITTING
    add_heading('10. Handling Overfitting & Underfitting', level=1)
    doc.add_paragraph("To mitigate overfitting and underfitting, several strategies were employed:")
    doc.add_paragraph("• Ensemble Approaches: For both anomaly detection and forecasting, ensemble methods were used. The union voting ensemble for anomalies and the Prophet+LightGBM ensemble for forecasting naturally reduce variance and prevent overfitting to specific noise patterns in the training data.", style='List Bullet')
    doc.add_paragraph("• SHAP Interpretability: SHAP values are used to ensure the model is making decisions based on sensible features (e.g., flagging an anomaly because of a massive surge in a specific service's usage, rather than an irrelevant background feature).", style='List Bullet')
    doc.add_paragraph("• Graceful Degradation: The system is designed to gracefully degrade if the SHAP library is unavailable or fails to compute, ensuring the core detection pipeline remains functional.", style='List Bullet')
    doc.add_paragraph("Because the models are primarily statistical (STL, Prophet) and tree-based (Isolation Forest, LightGBM), complex deep learning transfer learning or fine-tuning techniques were not required, reducing the risk of memorization.")

    # 11. SYSTEM ARCHITECTURE
    add_heading('11. System Architecture', level=1)
    doc.add_paragraph("The FinOps Intelligence Platform follows a decoupled, full-stack architecture designed for scalability and real-time processing.")
    doc.add_paragraph("• Frontend: A React SPA (Single Page Application) handles user interaction, state management, and rendering data visualizations.", style='List Bullet')
    doc.add_paragraph("• Backend API: A FastAPI backend acts as the central orchestration layer, exposing RESTful endpoints for data ingestion, querying, and triggering ML jobs.", style='List Bullet')
    doc.add_paragraph("• ML Pipeline: Background tasks (managed via FastAPI's BackgroundTasks) execute the computationally expensive anomaly detection and forecasting models asynchronously, preventing the API from blocking.", style='List Bullet')
    doc.add_paragraph("• Database Layer: SQLite configured in WAL mode provides robust concurrent read/write capabilities, allowing the frontend to poll for live updates while the ML pipeline writes detection results.", style='List Bullet')
    doc.add_paragraph("• Real-Time Streaming: During CSV uploads, the frontend utilizes an HTTP polling mechanism (500ms interval) to simulate a real-time data feed, fetching the latest detected anomalies as they are processed row-by-row by the backend.", style='List Bullet')
    doc.add_paragraph("The application features four main pages: the Command Center (main dashboard), Anomaly Watch (detailed incident review), Spend Forecasting (budget analysis), and Data Upload (real-time ingestion interface).")

    # 12. KEY DESIGN DECISIONS
    add_heading('12. Key Design Decisions', level=1)
    
    table2 = doc.add_table(rows=1, cols=2)
    table2.style = 'Table Grid'
    hdr_cells2 = table2.rows[0].cells
    hdr_cells2[0].text = 'Design Decision'
    hdr_cells2[1].text = 'Rationale'
    
    decisions = [
        ("SQLite with WAL mode", "Eliminates the need for a complex external database service while supporting the concurrent reads/writes necessary for real-time polling during uploads."),
        ("Union Ensemble Voting", "Maximizes anomaly recall. In FinOps, missing a massive cost spike is far more damaging than investigating a false positive."),
        ("Direct Multi-Step Forecasting", "Prevents compounding errors over time, providing more stable and reliable long-term budget forecasts compared to recursive methods."),
        ("FastAPI Background Tasks", "Ensures the main API thread remains highly responsive while offloading heavy ML computations to the background."),
        ("CSS Modules", "Provides locally scoped styling in React, preventing CSS conflicts across the 4 distinct dashboard pages."),
        ("Streaming via Polling", "Offers a robust, easy-to-implement real-time update mechanism without the infrastructure overhead of WebSocket/SSE connections."),
        ("SHAP as Optional", "Ensures the pipeline does not completely fail if SHAP computation times out or encounters dependency issues."),
        ("_free_port helper", "A backend utility that dynamically finds available ports, preventing port-collision crashes during development and deployment.")
    ]
    
    for dec, rat in decisions:
        row_cells = table2.add_row().cells
        row_cells[0].text = dec
        row_cells[1].text = rat

    # 13. RESULTS & DISCUSSION
    add_heading('13. Results & Discussion', level=1)
    doc.add_paragraph("The implemented FinOps Intelligence Platform successfully achieves its core objectives. During evaluation against the synthetic multi-cloud dataset:")
    doc.add_paragraph("• The anomaly detection ensemble successfully identified all 6 ground-truth anomalies (ANO-001 through ANO-006). The SHAP explanations accurately attributed the root cause to the correct CloudProvider, Service, and Region in each scenario.", style='List Bullet')
    doc.add_paragraph("• The forecasting engine generated accurate probabilistic bands, successfully providing early warnings when projected spend crossed the simulated budget thresholds.", style='List Bullet')
    doc.add_paragraph("• The Groq-powered Llama 3.3 chatbot effectively translated natural language queries into accurate SQLite queries, retrieving context-aware insights regarding cost breakdowns and anomalies.", style='List Bullet')
    doc.add_paragraph("• The real-time CSV upload architecture successfully processed 60,000+ rows of billing data, providing a live feed of anomalies to the frontend dashboard with minimal latency.")

    # 14. CHALLENGES FACED
    add_heading('14. Challenges Faced', level=1)
    doc.add_paragraph("Several technical challenges were encountered and overcome during development:")
    doc.add_paragraph("• Schema Normalization: Unifying the wildly different billing export formats of AWS (CUR), Azure, and GCP into a single, cohesive schema required careful data modeling.", style='List Bullet')
    doc.add_paragraph("• Precision vs. Recall Trade-off: Tuning the Isolation Forest and STL algorithms to find the sweet spot where legitimate seasonal traffic spikes were not flagged, while still catching subtle but expensive gradual drifts.", style='List Bullet')
    doc.add_paragraph("• Real-Time Architecture: Implementing a smooth real-time upload experience without WebSockets required fine-tuning the FastAPI endpoints and React polling intervals to avoid overwhelming the server.", style='List Bullet')
    doc.add_paragraph("• Explainability Overhead: Computing SHAP values for tree-based models on large datasets proved computationally intensive, necessitating background task offloading and making the SHAP calculation an optional/fallback step.")

    # 15. POTENTIAL IMPROVEMENTS
    add_heading('15. Potential Improvements', level=1)
    
    table3 = doc.add_table(rows=1, cols=2)
    table3.style = 'Table Grid'
    hdr_cells3 = table3.rows[0].cells
    hdr_cells3[0].text = 'Enhancement Area'
    hdr_cells3[1].text = 'Description'
    
    improvements = [
        ("Chat History Persistence", "Store conversation history in SQLite to provide the LLM with conversational context across multiple turns."),
        ("Hyperparameter Tuning", "Integrate Optuna to automatically tune LightGBM and Isolation Forest parameters based on historical data."),
        ("WebSocket/SSE Implementation", "Upgrade from HTTP polling to WebSockets or Server-Sent Events for more efficient, true real-time streaming."),
        ("Authentication & Authorization", "Implement JWT-based auth and Role-Based Access Control (RBAC) to restrict access to sensitive financial data."),
        ("Docker Containerization", "Provide Dockerfiles and docker-compose configurations to streamline deployment across different environments."),
        ("Multi-User Session Management", "Support concurrent users uploading and analyzing different billing datasets simultaneously."),
        ("Expanded Ingestion Support", "Add support for native Parquet and JSON cloud billing exports, rather than relying solely on CSV files.")
    ]
    
    for area, desc in improvements:
        row_cells = table3.add_row().cells
        row_cells[0].text = area
        row_cells[1].text = desc

    # 16. CONCLUSION
    add_heading('16. Conclusion', level=1)
    doc.add_paragraph("The FinOps Intelligence Platform demonstrates the immense value of applying modern AI techniques to the domain of cloud financial operations. By combining robust anomaly detection ensembles, probabilistic forecasting models, and large language models for natural language interaction, the platform transforms raw, incomprehensible cloud billing data into actionable business intelligence. The system successfully addresses the core challenges of multi-cloud cost management, proving that automated AI solutions can effectively detect billing spikes, forecast budget breaches, and provide clear root-cause explanations. This project provides a solid foundation for a scalable, enterprise-grade FinOps tool.")

    # 17. REFERENCES
    add_heading('17. References', level=1)
    doc.add_paragraph("1. Taylor, S. J., & Letham, B. (2018). Forecasting at Scale. The American Statistician, 72(1), 37-45. (Facebook Prophet)", style='List Bullet')
    doc.add_paragraph("2. Liu, F. T., Ting, K. M., & Zhou, Z. H. (2008). Isolation Forest. 2008 Eighth IEEE International Conference on Data Mining, 413-422.", style='List Bullet')
    doc.add_paragraph("3. Lundberg, S. M., & Lee, S. I. (2017). A Unified Approach to Interpreting Model Predictions. Advances in Neural Information Processing Systems (SHAP).", style='List Bullet')
    doc.add_paragraph("4. Ke, G., et al. (2017). LightGBM: A Highly Efficient Gradient Boosting Decision Tree. Advances in Neural Information Processing Systems.", style='List Bullet')
    doc.add_paragraph("5. Meta Llama 3 Model Card. Available at: https://llama.meta.com/llama3/", style='List Bullet')
    doc.add_paragraph("6. Groq API Documentation. Available at: https://console.groq.com/docs/", style='List Bullet')
    doc.add_paragraph("7. FastAPI Documentation. Available at: https://fastapi.tiangolo.com/", style='List Bullet')
    
    doc.save('FinOps_Intelligence_Platform_Report.docx')

if __name__ == "__main__":
    main()
